#!/usr/bin/env python3
"""Parse 9月10月账单.xlsx → settlement_spec_v4.json (3 Sept customers + 1 Oct)."""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

XLSX = r"E:\weixincache\xwechat_files\oceanfishly_3f62\msg\file\2026-10\9月10月账单.xlsx"
OUT = Path(r"d:\nodejs\newapi\new-api\.tmp_settlement_gen\settlement_spec_v4.json")

# 客户 sheet → user account (供应商仅对应关系，不单独建账号)
CUSTOMERS = [
    {
        "sheet": "客户1北京云天 9月",
        "username": "cust_bjyt",
        "display": "北京云天",
        "remark": "结算:北京云天9月↔供应商1金酷",
        "supplier": "供应商1金酷 9月",
        "month": "2026-09",
        "year": 2026,
        "month_num": 9,
    },
    {
        "sheet": "客户2深圳云天 9月",
        "username": "cust_szyt",
        "display": "深圳云天",
        "remark": "结算:深圳云天9月↔供应商3浙江天铁/相关",
        "supplier": "供应商3浙江天铁",
        "month": "2026-09",
        "year": 2026,
        "month_num": 9,
    },
    {
        "sheet": "客户3中联投 9月",
        "username": "cust_zlt",
        "display": "中联投",
        "remark": "结算:中联投9月↔供应商2曦腾等",
        "supplier": "供应商2曦腾 9月",
        "month": "2026-09",
        "year": 2026,
        "month_num": 9,
    },
    {
        "sheet": "客户3中联投10月",
        "username": "cust_zlt_oct",
        "display": "中联投-10月",
        "remark": "结算:中联投10月↔供应商曦腾10月",
        "supplier": "供应商曦腾10月",
        "month": "2026-10",
        "year": 2026,
        "month_num": 10,
    },
]


def parse_num(v):
    if v is None or v == "" or v == "/":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if s in ("", "/", "-", "—"):
        return None
    return float(s)


def excel_model_name(raw) -> str:
    parts = [p.strip() for p in str(raw).replace("\r", "\n").split("\n") if p.strip()]
    if not parts:
        return ""
    s = parts[0]
    for p in parts[1:]:
        s += p if p.startswith(("（", "(")) else (" " + p)
    while "  " in s:
        s = s.replace("  ", " ")
    return s.strip()


def is_video_model(name: str, qty_out, qty_cache) -> bool:
    n = name.lower()
    if "seedance" in n:
        return True
    if qty_out is None and qty_cache is None:
        return True
    return False


def parse_sheet(ws):
    items = []
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None or str(row[1]).strip() != "月":
            continue
        name = excel_model_name(row[0])
        if not name or name.startswith(("资源", "合计")):
            continue
        context = "" if row[2] is None or str(row[2]).strip() == "/" else str(row[2]).strip()
        qty_in = parse_num(row[3])
        qty_out = parse_num(row[4])
        qty_cache = parse_num(row[5])
        price_in = parse_num(row[6]) or 0.0
        price_out = parse_num(row[7]) or 0.0
        price_cache = parse_num(row[8]) or 0.0
        amount = parse_num(row[9])
        if amount is None or amount <= 0 or qty_in is None:
            continue
        video = is_video_model(name, qty_out, qty_cache)
        items.append(
            {
                "model": name,
                "context": context,
                "qty_in_m": float(qty_in),
                "qty_out_m": float(qty_out or 0),
                "qty_cache_m": float(qty_cache or 0),
                "price_in": float(price_in),
                "price_out": float(price_out),
                "price_cache": float(price_cache),
                "amount": float(amount),
                "is_video": video,
            }
        )
    return items


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    sheet_map = {s.strip(): s for s in wb.sheetnames}
    customers = []
    model_ratio = {}
    completion_ratio = {}
    cache_ratio = {}
    for meta in CUSTOMERS:
        raw = meta["sheet"]
        # fuzzy match sheet name (trailing spaces)
        sheet_name = sheet_map.get(raw) or sheet_map.get(raw.strip())
        if not sheet_name:
            for k, v in sheet_map.items():
                if k.replace(" ", "") == raw.replace(" ", ""):
                    sheet_name = v
                    break
        if not sheet_name:
            raise SystemExit(f"sheet not found: {raw!r} have={list(sheet_map)}")
        items = parse_sheet(wb[sheet_name])
        total = sum(i["amount"] for i in items)
        for it in items:
            if it["price_in"] > 0:
                model_ratio[it["model"]] = it["price_in"] / 2.0
            if it["price_in"] > 0 and it["price_out"] > 0:
                completion_ratio[it["model"]] = it["price_out"] / it["price_in"]
            if it["price_in"] > 0 and it["price_cache"] > 0:
                cache_ratio[it["model"]] = it["price_cache"] / it["price_in"]
        customers.append(
            {
                **{k: meta[k] for k in ("username", "display", "remark", "supplier", "month", "year", "month_num")},
                "sheet": sheet_name,
                "items": items,
                "total_amount": total,
            }
        )
        print(meta["username"], meta["month"], f"¥{total:,.2f}", f"lines={len(items)}")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    spec = {
        "password": "YxzSep2026!",
        "quota_per_unit": 500000.0,
        "video_price_range": [5.0, 20.0],
        "text_price_range": [1.0, 8.0],
        "model_ratio": model_ratio,
        "completion_ratio": completion_ratio,
        "cache_ratio": cache_ratio,
        "customers": customers,
        "source_xlsx": XLSX,
    }
    OUT.write_text(json.dumps(spec, ensure_ascii=False, indent=2), encoding="utf-8")
    print("wrote", OUT, "size", OUT.stat().st_size)
    print("grand_total", sum(c["total_amount"] for c in customers))


if __name__ == "__main__":
    main()
